from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
# pyrefly: ignore [missing-import]
from docx.oxml import OxmlElement
# pyrefly: ignore [missing-import]
from docx.oxml.ns import qn
# pyrefly: ignore [missing-import]
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Tedarikci_Risk_Degerlendirmesi_Raporu_REVIZE.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(88, 88, 88)
GOLD = RGBColor(127, 104, 49)
BLACK = RGBColor(0, 0, 0)
LIGHT_GRAY = "F2F4F7"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = str(h)
        shade_cell(hdr[i], LIGHT_GRAY)
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, size=font_size, bold=True, color=BLACK)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    set_run_font(r, size=font_size, color=BLACK)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph()
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 11),
                     color=BLUE if level in (1, 2) else DARK_BLUE, bold=True)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    return p


def add_para(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=9, italic=True, color=MUTED)


def add_image(doc, filename, caption, width=5.9):
    path = ROOT / filename
    if not path.exists():
        add_para(doc, f"[Görsel bulunamadı: {filename}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def add_toc(doc):
    entries = [
        "1. İçindekiler",
        "2. Giriş",
        "3. Literatür Özeti",
        "4. Teorik Arka Plan",
        "5. Veri Seti Açıklaması",
        "6. Yöntem ve Uygulama",
        "7. Sonuçlar ve Yorum",
        "8. Sonuç ve Değerlendirme",
        "9. Kaynakça",
        "10. Ekler",
    ]
    for entry in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(entry)
        set_run_font(r, size=11)


def fmt_float(x):
    return f"{float(x):.4f}"


def model_rows(file_name):
    df = pd.read_csv(ROOT / file_name)
    model_name_map = {
        "Logistic Regression": "Lojistik Regresyon",
        "Decision Tree": "Karar Ağacı",
        "Random Forest": "Random Forest",
        "K-Nearest Neighbors": "K-En Yakın Komşu",
    }
    return [
        [model_name_map.get(r["Model"], r["Model"]), fmt_float(r["Accuracy"]), fmt_float(r["Precision_Macro"]),
         fmt_float(r["Recall_Macro"]), fmt_float(r["F1_Macro"])]
        for _, r in df.iterrows()
    ]


def appendix_inventory_rows():
    rows = [
        ("supplier_risk.csv", "Ham veri seti."),
        ("supplier_risk_scored.csv", "Sabit eşik risk skorları eklenmiş veri seti."),
        ("supplier_risk_scored_quantile.csv", "Kuantil sınıfları eklenmiş nihai veri seti."),
        ("eda_supplier_risk.py", "Keşifçi veri analizi kodu."),
        ("supplier_scoring.py", "Sabit eşik risk skorlama kodu."),
        ("supplier_scoring_quantile.py", "Kuantil eşik sınıflandırma kodu."),
        ("ml_modeling.py", "Pipeline tabanlı modelleme kodu."),
        ("advanced_stats.py", "VIF ve R² hesaplama kodu."),
        ("quantile_thresholds.csv", "Kuantil eşik değerleri."),
        ("vif_results.csv", "VIF analizi çıktısı."),
        ("r2_results.txt", "R² analizi çıktısı."),
        ("model_comparison_priority_fixed_without_year.csv", "Ana model, öncelik ağırlıklı skor, Year hariç model karşılaştırması."),
        ("model_comparison_equal_fixed_without_year.csv", "Ana model, eşit ağırlıklı skor, Year hariç model karşılaştırması."),
        ("model_comparison_priority_quantile_without_year.csv", "Yan model, öncelik ağırlıklı skor, Year hariç model karşılaştırması."),
        ("model_comparison_equal_quantile_without_year.csv", "Yan model, eşit ağırlıklı skor, Year hariç model karşılaştırması."),
        ("model_comparison_priority_fixed_with_year.csv", "Year dahil karşılaştırmalı model çıktısı."),
        ("model_comparison_equal_fixed_with_year.csv", "Year dahil karşılaştırmalı model çıktısı."),
        ("model_comparison_priority_quantile_with_year.csv", "Year dahil karşılaştırmalı model çıktısı."),
        ("model_comparison_equal_quantile_with_year.csv", "Year dahil karşılaştırmalı model çıktısı."),
        ("risk_category_distribution.png", "Orijinal risk kategorisi dağılım görselleştirmesi."),
        ("year_distribution.png", "Yıl dağılımı görselleştirmesi."),
        ("correlation_heatmap.png", "Korelasyon ısı haritası."),
        ("risk_class_comparison.png", "Sabit eşik risk sınıfları karşılaştırma grafiği."),
        ("threshold_comparison.png", "Sabit eşik ve kuantil eşik karşılaştırma grafiği."),
        ("confusion_matrices_priority_fixed_without_year.png", "Ana model öncelik ağırlıklı skor confusion matrix çıktısı."),
        ("confusion_matrices_equal_fixed_without_year.png", "Ana model eşit ağırlıklı skor confusion matrix çıktısı."),
        ("confusion_matrices_priority_quantile_without_year.png", "Yan model öncelik ağırlıklı skor confusion matrix çıktısı."),
        ("confusion_matrices_equal_quantile_without_year.png", "Yan model eşit ağırlıklı skor confusion matrix çıktısı."),
        ("rf_feature_importance_priority_fixed_without_year.png", "Ana model öncelik ağırlıklı skor Random Forest değişken önemleri."),
        ("rf_feature_importance_equal_fixed_without_year.png", "Ana model eşit ağırlıklı skor Random Forest değişken önemleri."),
        ("rf_feature_importance_priority_quantile_without_year.png", "Yan model öncelik ağırlıklı skor Random Forest değişken önemleri."),
        ("rf_feature_importance_equal_quantile_without_year.png", "Yan model eşit ağırlıklı skor Random Forest değişken önemleri."),
    ]
    return [[name, desc] for name, desc in rows if (ROOT / name).exists()]


def add_code_appendix(doc, file_name):
    add_heading(doc, file_name, 2)
    text = (ROOT / file_name).read_text(encoding="utf-8")
    lines = text.splitlines()
    for i in range(0, len(lines), 38):
        chunk = "\n".join(lines[i:i + 38])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(chunk)
        set_run_font(r, name="Courier New", size=7.5)


def build():
    df = pd.read_csv(ROOT / "supplier_risk.csv")
    scored = pd.read_csv(ROOT / "supplier_risk_scored_quantile.csv")
    vif = pd.read_csv(ROOT / "vif_results.csv")
    quantiles = pd.read_csv(ROOT / "quantile_thresholds.csv")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    for s_name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 12, DARK_BLUE)]:
        style = styles[s_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    # Cover
    logo = ROOT / "ytu_logo.png"
    if logo.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(logo), width=Inches(1.25))
    for text, size, bold, color, after in [
        ("YILDIZ TEKNİK ÜNİVERSİTESİ", 18, True, BLACK, 2),
        ("İstatistik Bölümü", 13, False, MUTED, 34),
        ("VERİ MADENCİLİĞİNE GİRİŞ DÖNEM ÖDEVİ", 18, True, BLUE, 10),
        ("Tedarikçi Risk Değerlendirmesi: Veri Madenciliği Yaklaşımı", 15, True, BLACK, 34),
        ("Öğrenci Adı Soyadı: ........................................", 11, False, BLACK, 3),
        ("Öğrenci No: ................................................", 11, False, BLACK, 18),
        ("2025-2026 Bahar Dönemi", 12, True, GOLD, 0),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    doc.add_page_break()

    add_heading(doc, "1. İçindekiler", 1)
    add_toc(doc)
    doc.add_page_break()

    add_heading(doc, "2. Giriş", 1)
    add_heading(doc, "2.1 Konunun Tanıtımı", 2)
    add_para(doc, "Tedarik zinciri yönetiminde tedarikçi risk değerlendirmesi, işletmelerin tedarikçilerini finansal istikrar, teslimat performansı, kalite uyumluluğu, düzenleyici uygunluk, sürdürülebilirlik, geçmiş risk seviyesi ve olay sayısı gibi ölçütlere göre sistematik biçimde değerlendirmesini sağlar. Bu süreç yalnızca satın alma kararlarını değil, üretim sürekliliğini, stok planlamasını ve operasyonel dayanıklılığı da doğrudan etkiler.")
    add_para(doc, "Veri madenciliği teknikleri, bu değerlendirme sürecinin öznel yorumlardan uzaklaştırılarak ölçülebilir ve tekrar üretilebilir bir yapıya taşınmasına yardımcı olur. Bu projede tedarikçi performans metrikleri önce keşifçi veri analiziyle incelenmiş, ardından kompozit risk skorlarına ve makine öğrenmesi sınıflandırma modellerine dönüştürülmüştür.")
    add_heading(doc, "2.2 Konunun Önemi", 2)
    add_para(doc, "Küreselleşen tedarik zincirlerinde tek bir tedarikçide yaşanan aksama, üretim hattı, teslimat süreleri ve müşteri memnuniyeti üzerinde zincirleme etki yaratabilir. COVID-19 döneminde görünür hale gelen kırılganlıklar, risk değerlendirmesinin yalnızca kriz anında değil sürekli izlenmesi gereken stratejik bir faaliyet olduğunu göstermiştir.")
    add_heading(doc, "2.3 Ödevin Amacı ve Hedefleri", 2)
    add_para(doc, "Bu projenin amacı Kaggle kaynaklı Supplier Risk Assessment veri seti üzerinde keşifçi veri analizi yapmak, mevcut Risk_Category değişkenini metodolojik olarak değerlendirmek, bağımsız bir kompozit risk skorlama sistemi oluşturmak ve sınıflandırma algoritmalarıyla tedarikçilerin risk sınıflarını tahmin etmektir.")
    add_bullets(doc, [
        "Veri kalitesini, değişken dağılımlarını ve korelasyon yapısını incelemek.",
        "Risk skorunu oluşturan değişkenleri gerekçeli biçimde seçmek.",
        "Eşit ağırlıklı ve öncelik ağırlıklı iki farklı skor üretmek.",
        "Sabit eşik ve kuantil eşik sınıflandırmalarını karşılaştırmak.",
        "Lojistik Regresyon, Karar Ağacı, Random Forest ve KNN modellerini test etmek.",
        "VIF ve R² analizleriyle model yapısının istatistiksel tutarlılığını değerlendirmek.",
    ])
    add_heading(doc, "2.4 Kapsam", 2)
    add_para(doc, "Çalışma; EDA, özellik mühendisliği, risk skorlama, sınıflandırma modellemesi, model karşılaştırması ve istatistiksel geçerlilik kontrollerini kapsamaktadır. Zaman serisi modellemesi, derin öğrenme, gerçek zamanlı izleme sistemi ve dış veri entegrasyonu bu ödevin kapsamı dışında bırakılmıştır.")

    add_heading(doc, "3. Literatür Özeti", 1)
    add_para(doc, "Tedarik zinciri risk yönetimi literatürü, firmaların kesinti olasılığını azaltmak ve kesinti gerçekleştiğinde etkisini sınırlamak için proaktif karar mekanizmalarına ihtiyaç duyduğunu vurgular. Tang (2006), tedarik zinciri risk yönetimini hem operasyonel uygulamalar hem de araştırma yaklaşımları açısından kapsamlı biçimde ele alır ve risk azaltma stratejilerinin yalnızca maliyet değil süreklilik perspektifiyle değerlendirilmesi gerektiğini belirtir.")
    add_para(doc, "Tedarikçi seçimi ve değerlendirmesi genellikle çok kriterli karar verme problemi olarak ele alınır. Ho, Xu ve Dey (2010), tedarikçi değerlendirmede tek ölçütlü maliyet yaklaşımının yetersiz kaldığını; kalite, teslimat, finansal yapı ve sürdürülebilirlik gibi birden çok kriterin birlikte değerlendirilmesi gerektiğini gösteren geniş bir literatür taraması sunar.")
    add_para(doc, "Makine öğrenmesi yaklaşımları, tedarik zinciri risk yönetiminde risk tahmini ve erken uyarı üretme açısından giderek daha fazla kullanılmaktadır. Baryannis, Dani ve Antoniou (2019), tedarik zinciri risk tahmininde performans ile yorumlanabilirlik arasındaki dengenin önemli olduğunu vurgular. Bu proje bu çizgiyle uyumlu olarak hem yüksek performanslı hem de formülü açıklanabilir bir skor-model yapısı kurmuştur.")

    add_heading(doc, "4. Teorik Arka Plan", 1)
    add_heading(doc, "4.1 Temel Kavramlar", 2)
    add_bullets(doc, [
        "Sınıflandırma: Gözlemleri önceden tanımlanmış sınıflara atama işlemidir. Bu çalışmada Low Risk, Medium Risk ve High Risk sınıfları kullanılmıştır.",
        "Özellik mühendisliği: Ham değişkenlerden modelleme için anlamlı risk bileşenleri üretme sürecidir. Performans değişkenleri ters çevrilerek risk göstergesine dönüştürülmüştür.",
        "Normalizasyon: Değişkenleri ortak ölçeğe taşıma işlemidir. Skor üretiminde MinMaxScaler, model pipeline içinde ise Lojistik Regresyon ve K-En Yakın Komşu için StandardScaler kullanılmıştır.",
        "Kuantil eşik: Gözlemleri veri dağılımına göre alt, orta ve üst dilimlere ayıran eşikleme yaklaşımıdır.",
    ])
    add_heading(doc, "4.2 Kullanılan Algoritmalar", 2)
    add_para(doc, "Lojistik Regresyon, sınıflar arasındaki ayrımı doğrusal karar sınırlarıyla modelleyen yorumlanabilir bir sınıflandırma algoritmasıdır. Risk sınıfları doğrusal/ağırlıklı skor mantığıyla üretildiği için bu modelin yüksek performans göstermesi beklenir.")
    add_para(doc, "Karar Ağacı, veriyi ardışık karar kurallarıyla dallandırır. Ölçeklendirme gerektirmez ve yorumlanabilirlik avantajı sunar; ancak sınıf sınırları doğrusal veya çok hassas olduğunda performansı dalgalanabilir.")
    add_para(doc, "Random Forest, çok sayıda karar ağacının birlikte çalıştığı topluluk yöntemidir. Ağaçların ortalaması sayesinde tekil karar ağacına göre daha kararlı sonuçlar üretir ve değişken önem sıralaması sağlar.")
    add_para(doc, "K-En Yakın Komşu, yeni gözlemi eğitim setindeki en yakın komşularına göre sınıflandırır. Mesafe tabanlı olduğu için ölçek farklılıklarından güçlü biçimde etkilenir; bu nedenle StandardScaler ile pipeline içinde kullanılmıştır.")
    add_heading(doc, "4.3 Değerlendirme Metrikleri", 2)
    add_bullets(doc, [
        "Accuracy: Toplam doğru tahminlerin tüm tahminlere oranıdır.",
        "Precision Macro: Her sınıf için kesinlik değerlerinin ortalamasıdır.",
        "Recall Macro: Her sınıf için duyarlılık değerlerinin ortalamasıdır.",
        "F1 Macro: Precision ve recall değerlerinin harmonik ortalamasıdır; sınıflar arası dengeyi daha iyi yansıtır.",
        "Confusion Matrix: Gerçek ve tahmin edilen sınıfları çapraz biçimde gösterir.",
        "VIF: Bağımsız değişkenler arasında çoklu doğrusal bağlantı olup olmadığını ölçer.",
        "R²: Sürekli risk skorlarının seçilen değişkenler tarafından ne ölçüde açıklandığını gösterir.",
    ])

    add_heading(doc, "5. Veri Seti Açıklaması", 1)
    add_heading(doc, "5.1 Genel Bilgiler", 2)
    add_table(doc, ["Özellik", "Değer"], [
        ["Veri seti", "Supplier Risk Assessment Dataset"],
        ["Kaynak", "Kaggle"],
        ["Gözlem sayısı", f"{df.shape[0]}"],
        ["Değişken sayısı", f"{df.shape[1]}"],
        ["Benzersiz tedarikçi", f"{df['Supplier_ID'].nunique()}"],
        ["Yıl aralığı", f"{int(df['Year'].min())}-{int(df['Year'].max())}"],
        ["Eksik değer", f"{int(df.isna().sum().sum())}"],
        ["Tekrarlı satır", f"{int(df.duplicated().sum())}"],
    ], widths=[2.2, 4.1])
    add_para(doc, "Veri setine ilişkin temel bilgiler doğrudan analiz edilen CSV dosyası üzerinden doğrulanmıştır.")
    add_para(doc, "Eksik değer ve tekrarlı satır kontrollerinin ardından sayısal değişkenler kutu grafikleriyle incelenmiştir. Skor değişkenleri beklenen aralıklarda yer aldığı ve gözlemler tedarikçi performans farklılıklarını temsil ettiği için ayrıca aykırı değer silme işlemi uygulanmamıştır.")
    add_heading(doc, "5.2 Değişkenler", 2)
    add_table(doc, ["Değişken", "Tür", "Açıklama"], [
        ["Supplier_ID", "Tamsayı", "Tedarikçi kimlik numarası."],
        ["Year", "Tamsayı", "Gözlemin ait olduğu yıl."],
        ["Financial_Stability_Score", "Ondalıklı", "Finansal istikrar; yüksek değer olumlu."],
        ["Delivery_Performance_Score", "Ondalıklı", "Teslimat performansı; yüksek değer olumlu."],
        ["Quality_Compliance_Score", "Ondalıklı", "Kalite uyumluluğu; yüksek değer olumlu."],
        ["Regulatory_Adherence_Score", "Ondalıklı", "Düzenleyici uyum; yüksek değer olumlu."],
        ["Sustainability_Score", "Ondalıklı", "Sürdürülebilirlik performansı; yüksek değer olumlu."],
        ["Past_Risk_Level", "Ondalıklı", "Geçmiş risk seviyesi; yüksek değer riskli."],
        ["ERP_Transactions", "Tamsayı", "ERP işlem sayısı; ana risk skoruna dahil edilmemiştir."],
        ["Incidents_Count", "Tamsayı", "Olay/sorun sayısı; yüksek değer riskli."],
        ["MCDM_Score", "Ondalıklı", "Önceden hesaplanmış kompozit skor; çift sayım riski nedeniyle kullanılmamıştır."],
        ["Risk_Category", "Kategorik", "Orijinal kategori; hedef değişken olarak kullanılmamıştır."],
    ], widths=[2.25, 1.25, 3.0], font_size=8.3)
    add_heading(doc, "5.3 Keşifçi Veri Analizi Bulguları", 2)
    add_para(doc, f"Orijinal Risk_Category dağılımında {df['Risk_Category'].value_counts().to_dict()} sonucu elde edilmiştir. Veri setinde Low kategorisi gözlenmemekte, yalnızca Medium ve High sınıfları bulunmaktadır.")
    add_para(doc, "Orijinal High kategorisinin daha yüksek performans skorlarıyla ilişkili olduğu görülmüştür. Örneğin High grubunun ortalama Financial_Stability_Score değeri 0.769 iken Medium grubunda 0.531'dir. Bu nedenle Risk_Category değişkeni doğrudan risk hedefi olarak kullanılmamıştır.")
    add_image(doc, "risk_category_distribution.png", "Şekil 1. Orijinal Risk_Category dağılımı.", width=5.2)
    add_image(doc, "year_distribution.png", "Şekil 2. Gözlemlerin yıllara göre dağılımı.", width=5.2)
    add_image(doc, "correlation_heatmap.png", "Şekil 3. Sayısal değişkenler için korelasyon ısı haritası.", width=5.8)
    add_image(doc, "boxplot_Financial_Stability_Score.png", "Şekil 4. Financial_Stability_Score değişkeninin Risk_Category bazında kutu grafiği.", width=5.2)
    add_image(doc, "boxplot_Delivery_Performance_Score.png", "Şekil 5. Delivery_Performance_Score değişkeninin Risk_Category bazında kutu grafiği.", width=5.2)

    add_heading(doc, "6. Yöntem ve Uygulama", 1)
    add_heading(doc, "6.1 Genel Yaklaşım", 2)
    add_numbered(doc, [
        "Keşifçi veri analizi ile veri kalitesi, dağılımlar ve korelasyon yapısı incelenmiştir.",
        "Risk bileşenleri oluşturularak eşit ve öncelik ağırlıklı iki kompozit skor tasarlanmıştır.",
        "Sabit eşik ve kuantil eşik yöntemleriyle risk sınıfları üretilmiştir.",
        "Dört sınıflandırma algoritması sklearn Pipeline yapısıyla eğitilip karşılaştırılmıştır.",
    ])
    add_heading(doc, "6.2 Risk Skorlama Sistemi", 2)
    add_para(doc, "Finansal istikrar, teslimat performansı, kalite uyumluluğu, düzenleyici uyum ve sürdürülebilirlik değişkenleri yüksek olduğunda iyi performansı gösterdiği için önce MinMaxScaler ile 0-1 aralığına alınmış, ardından 1 - normalize_değer dönüşümüyle risk bileşenine çevrilmiştir. Past_Risk_Level ve Incidents_Count yüksek olduğunda doğrudan risk göstergesi olduğu için normalize değerleri doğrudan kullanılmıştır.")
    add_table(doc, ["Bileşen", "Dönüşüm", "Gerekçe"], [
        ["Financial_Stability_Score", "1 - normalized", "Yüksek finansal istikrar düşük risktir."],
        ["Delivery_Performance_Score", "1 - normalized", "Yüksek teslimat performansı düşük risktir."],
        ["Quality_Compliance_Score", "1 - normalized", "Yüksek kalite uyumu düşük risktir."],
        ["Regulatory_Adherence_Score", "1 - normalized", "Yüksek uyum düşük risktir."],
        ["Sustainability_Score", "1 - normalized", "Yüksek sürdürülebilirlik düşük risktir."],
        ["Past_Risk_Level", "normalized", "Yüksek geçmiş risk doğrudan risktir."],
        ["Incidents_Count", "normalized", "Yüksek olay sayısı doğrudan risktir."],
    ], widths=[2.2, 1.6, 2.7], font_size=8.5)
    add_para(doc, "Eşit ağırlıklı skorda yedi risk bileşeninin aritmetik ortalaması alınmıştır. Öncelik ağırlıklı skorda finansal istikrar ve teslimat performansı daha yüksek ağırlık taşımıştır: 0.27 financial + 0.21 delivery + 0.13 quality + 0.08 regulatory + 0.11 sustainability + 0.10 past risk + 0.10 incident risk.")
    add_heading(doc, "6.3 Sınıflandırma Eşikleri", 2)
    add_table(doc, ["Yöntem", "Low Risk", "Medium Risk", "High Risk"], [
        ["Sabit eşik", "0.00-0.33", "0.33-0.66", "0.66-1.00"],
        ["Kuantil eşik", "Alt %33", "Orta %33", "Üst %34"],
    ], widths=[1.7, 1.6, 1.6, 1.6])
    add_table(doc, ["Skorlama yöntemi", "Low/Medium eşiği", "Medium/High eşiği"], [
        [r["Scoring_Method"], fmt_float(r["Low_Medium_Threshold"]), fmt_float(r["Medium_High_Threshold"])]
        for _, r in quantiles.iterrows()
    ], widths=[2.4, 2.0, 2.0])
    add_para(doc, "Sabit eşik yöntemi, 0-1 aralığında üretilen risk skorunu teorik ve yorumlanabilir sınırlarla üç risk düzeyine ayırdığı için çalışmanın ana sınıflandırma yaklaşımı olarak ele alınmıştır. Kuantil eşik yöntemi ise sabit eşikte gözlenen sınıf yoğunlaşmasının sonuçlara etkisini incelemek ve sınıf dağılımına duyarlılığı test etmek amacıyla yan/karşılaştırmalı yaklaşım olarak kullanılmıştır.")
    add_para(doc, "Sabit eşikte gözlemlerin Medium Risk sınıfında yoğunlaşması, bu yaklaşımın teorik ana eşik yapısını geçersiz kılmamaktadır; yalnızca veri setindeki risk skorlarının orta bölgede kümelendiğini göstermektedir. Kuantil eşik ise yaklaşık 594/594/612 dağılımıyla sınıf dengesini iyileştirir, ancak risk düzeylerini mutlak skor aralıklarından çok veri seti içindeki göreli sıralama üzerinden tanımlar.")
    add_image(doc, "risk_class_comparison.png", "Şekil 6. Eşit ve öncelik ağırlıklı sabit eşik risk sınıfları.", width=5.9)
    add_image(doc, "threshold_comparison.png", "Şekil 7. Sabit eşik ve kuantil eşik sınıflandırmalarının karşılaştırması.", width=5.9)
    add_heading(doc, "6.4 Makine Öğrenmesi Pipeline Yapısı", 2)
    add_para(doc, "Modelleme aşamasında train-test ayrımı %80/%20 oranında yapılmış, random_state=42 seçilmiş ve hedef sınıf oranlarının korunması için stratify=y kullanılmıştır. Lojistik Regresyon ve K-En Yakın Komşu, StandardScaler içeren sklearn Pipeline yapısıyla kurulmuştur. Karar Ağacı ve Random Forest ölçeğe duyarlı olmadığından doğrudan kullanılmıştır.")
    add_table(doc, ["Model", "Ön işleme", "Gerekçe"], [
        ["Lojistik Regresyon", "StandardScaler + model", "Doğrusal modelin optimizasyonu ölçekten etkilenir."],
        ["K-En Yakın Komşu", "StandardScaler + model", "Mesafe tabanlı olduğu için ortak ölçek gerekir."],
        ["Karar Ağacı", "Ölçeklendirme yok", "Bölünme kuralları ölçekten bağımsızdır."],
        ["Random Forest", "Ölçeklendirme yok", "Ağaç tabanlı topluluk yöntemi ölçekten bağımsızdır."],
    ], widths=[2.0, 2.0, 2.5])
    add_heading(doc, "6.5 Modelleme Ayarları ve Hiperparametreler", 2)
    add_para(doc, "Modelleme aşamasında hiperparametre optimizasyonu yapılmamıştır. Varsayılan parametreler kullanılmış, yalnızca aşağıdaki açık ayarlar sabitlenmiştir.")
    add_table(doc, ["Bileşen", "Ayar"], [
        ["Train-test split", "80/20"],
        ["random_state", "42"],
        ["stratify", "y"],
        ["Lojistik Regresyon", "StandardScaler + LogisticRegression(max_iter=1000, random_state=42)"],
        ["Karar Ağacı", "DecisionTreeClassifier(random_state=42)"],
        ["Random Forest", "RandomForestClassifier(random_state=42)"],
        ["K-En Yakın Komşu", "StandardScaler + KNeighborsClassifier(default n_neighbors=5)"],
    ], widths=[2.1, 4.2], font_size=8.7)

    add_heading(doc, "7. Sonuçlar ve Yorum", 1)
    add_heading(doc, "7.1 İstatistiksel Geçerlilik Testleri", 2)
    add_heading(doc, "7.1.1 VIF Analizi", 3)
    add_table(doc, ["Özellik", "VIF"], [[r["Feature"], f"{float(r['VIF']):.3f}"] for _, r in vif.iterrows()], widths=[4.5, 1.2])
    add_para(doc, "Tüm VIF değerleri 1.00 civarındadır. Genel kabul olarak VIF < 5 olması çoklu doğrusal bağlantı sorununun ciddi olmadığını gösterir. Bu nedenle seçilen yedi değişkenin risk skorlama sisteminde birlikte kullanılması istatistiksel olarak uygundur.")
    add_heading(doc, "7.1.2 R² Analizi", 3)
    add_para(doc, "R² değerinin 1.0000 çıkması beklenen bir sonuçtur; çünkü Risk_Score_Equal ve Risk_Score_Priority doğrudan seçilen yedi değişkenin matematiksel birleşimiyle oluşturulmuştur. Bu nedenle bu analiz bağımsız bir tahmin başarısından ziyade, oluşturulan skorların seçilen değişkenlerle birebir ilişkili olduğunu göstermektedir.")

    add_heading(doc, "7.2 Ana Model Sonuçları: Sabit Eşik", 2)
    add_heading(doc, "7.2.1 Priority Fixed - Year Olmadan", 3)
    add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "F1"], model_rows("model_comparison_priority_fixed_without_year.csv"), widths=[2.2, 1, 1, 1, 1])
    add_heading(doc, "7.2.2 Equal Fixed - Year Olmadan", 3)
    add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "F1"], model_rows("model_comparison_equal_fixed_without_year.csv"), widths=[2.2, 1, 1, 1, 1])
    add_para(doc, "Sabit eşik yaklaşımında sınıflar dengesizdir. Medium Risk sınıfında yoğunlaşma oluştuğu için bazı modeller uç sınıfları yakalamakta zorlanmıştır. Ancak bu yoğunlaşma, sabit eşik yönteminin çalışmanın ana teorik sınıflandırma yaklaşımı olarak kullanılmasını geçersiz kılmaz; çünkü bu yöntem 0-1 risk skoru aralığını mutlak ve yorumlanabilir risk düzeylerine ayırmaktadır.")
    add_para(doc, "Lojistik Regresyon modelinin yüksek performansı, hedef risk sınıflarının bu çalışmada oluşturulan doğrusal/ağırlıklı kompozit risk skoru mantığına dayanmasıyla açıklanabilir. Bu nedenle sonuçlar, dışsal ve bağımsız bir gerçek dünya risk etiketini tahmin etmekten çok, geliştirilen risk skorlama sisteminin makine öğrenmesi modelleri tarafından ne ölçüde öğrenilebilir olduğunu göstermektedir.")

    add_heading(doc, "7.3 Yan Model Sonuçları: Kuantil Eşik", 2)
    add_heading(doc, "7.3.1 Priority Quantile - Year Olmadan", 3)
    add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "F1"], model_rows("model_comparison_priority_quantile_without_year.csv"), widths=[2.2, 1, 1, 1, 1])
    add_heading(doc, "7.3.2 Equal Quantile - Year Olmadan", 3)
    add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "F1"], model_rows("model_comparison_equal_quantile_without_year.csv"), widths=[2.2, 1, 1, 1, 1])
    add_para(doc, "Kuantil eşik yaklaşımı sınıfları dengeli hale getirdiği için Random Forest ve K-En Yakın Komşu gibi modellerin performansı daha kararlı görünmektedir. Bu yaklaşım, ana modelin yerini almak için değil, sabit eşikteki sınıf yoğunlaşmasının model sonuçlarını nasıl etkilediğini karşılaştırmalı olarak incelemek için kullanılmıştır. Kuantil eşik risk sınıflarını veri seti içindeki göreli sıralamaya göre tanımladığı için sonuçlar mutlak risk aralıklarından farklı yorumlanmalıdır.")
    add_para(doc, "Lojistik Regresyon modelinin kuantil eşikte de yüksek performans göstermesi, hedef sınıfların yine oluşturulan kompozit risk skoru mantığına dayanmasıyla uyumludur. Bu bulgu, modelin bağımsız bir dış risk etiketini keşfettiği anlamına değil, tasarlanan skor sisteminin sınıflandırma modelleri tarafından tutarlı biçimde öğrenilebildiğine işaret eder.")
    add_image(doc, "confusion_matrices_priority_fixed_without_year.png", "Şekil 8. Priority Fixed hedefi için confusion matrix sonuçları.", width=6.0)
    add_image(doc, "confusion_matrices_equal_fixed_without_year.png", "Şekil 9. Equal Fixed hedefi için confusion matrix sonuçları.", width=6.0)
    add_image(doc, "confusion_matrices_priority_quantile_without_year.png", "Şekil 10. Priority Quantile hedefi için confusion matrix sonuçları.", width=6.0)
    add_image(doc, "rf_feature_importance_priority_fixed_without_year.png", "Şekil 11. Priority Fixed Random Forest değişken önemleri.", width=5.9)
    add_image(doc, "rf_feature_importance_equal_quantile_without_year.png", "Şekil 12. Equal Quantile Random Forest değişken önemleri.", width=5.9)

    add_heading(doc, "7.4 Ana Model ve Yan Model Karşılaştırması", 2)
    add_table(doc, ["Kriter", "Sabit Eşik", "Kuantil Eşik"], [
        ["Sınıf dengesi", "Medium sınıfına yığılma vardır.", "Yaklaşık dengeli dağılım üretir."],
        ["Yorumlanabilirlik", "0-1 skor aralığında sezgisel sınırlar sunar.", "Veri seti içi göreli sıralamayı öne çıkarır."],
        ["En iyi model", "Lojistik Regresyon", "Lojistik Regresyon"],
        ["Ağaç modelleri", "Dengesiz sınıflardan etkilenir.", "Dengeli sınıflarda daha kararlı sonuç verir."],
        ["Metodolojik rol", "Ana model olarak sunulmuştur.", "Yan/karşılaştırmalı model olarak sunulmuştur."],
    ], widths=[1.7, 2.4, 2.4], font_size=8.8)
    add_para(doc, "Bu karşılaştırmada sabit eşik, çalışmanın ana modeli olarak korunmuştur; çünkü risk skorunun teorik 0-1 aralığını Low, Medium ve High Risk düzeylerine doğrudan bağlar. Kuantil eşik ise sabit eşikte oluşan Medium Risk yoğunlaşmasına karşı duyarlılık analizi niteliğindedir. Dolayısıyla kuantil yaklaşım sınıf dengesini artırsa da, proje metodolojisinde ana modelin yerine geçirilmemiştir.")
    add_heading(doc, "7.5 Year Değişkeninin Etkisi", 2)
    add_table(doc, ["Hedef", "Model", "Year Yok F1", "Year Var F1", "Fark"], [
        ["Priority Fixed", "Lojistik Regresyon", "0.9644", "0.9573", "-0.0071"],
        ["Priority Fixed", "Random Forest", "0.6948", "0.7095", "+0.0147"],
        ["Equal Fixed", "Lojistik Regresyon", "0.9626", "0.9626", "0.0000"],
        ["Equal Fixed", "Random Forest", "0.4683", "0.4412", "-0.0270"],
        ["Priority Quantile", "Lojistik Regresyon", "0.9944", "0.9944", "0.0000"],
        ["Priority Quantile", "Random Forest", "0.8559", "0.8569", "+0.0010"],
        ["Equal Quantile", "Lojistik Regresyon", "1.0000", "1.0000", "0.0000"],
        ["Equal Quantile", "Random Forest", "0.8579", "0.8367", "-0.0212"],
    ], widths=[1.7, 2.0, 1.0, 1.0, 0.8], font_size=8.5)
    add_para(doc, "Year değişkeni modele anlamlı ve tutarlı bir katkı sağlamamıştır. Risk skorları performans metriklerinden türetildiği için yıl bilgisi çoğu modelde yalnızca bağlamsal/gürültü niteliğinde kalmıştır. Bu nedenle nihai yorumlarda Year hariç model daha sade ve genellenebilir kabul edilmiştir.")

    add_heading(doc, "8. Sonuç ve Değerlendirme", 1)
    add_para(doc, "Bu projede 1800 gözlem ve 12 değişkenden oluşan Supplier Risk Assessment veri seti kullanılarak uçtan uca bir veri madenciliği çalışması yürütülmüştür. EDA sonucunda orijinal Risk_Category değişkeninin risk hedefi olarak doğrudan kullanılamayacağı görülmüş, bunun yerine seçilen yedi değişken üzerinden iki farklı kompozit risk skoru üretilmiştir.")
    add_para(doc, "Sabit eşik yaklaşımı çalışmanın ana teorik sınıflandırma yaklaşımıdır; çünkü 0-1 risk skoru aralığını yorumlanabilir Low, Medium ve High Risk düzeylerine ayırır. Kuantil eşik yaklaşımı ise sınıf dağılımına duyarlılığı ve model sonuçlarının robustluğunu değerlendirmek için kullanılan yan/karşılaştırmalı yaklaşımdır.")
    add_para(doc, "Modelleme aşamasında en başarılı algoritma Lojistik Regresyon olmuştur. Bu sonuç, hedef sınıfların çalışmada oluşturulan doğrusal/ağırlıklı risk skorundan türetilmesiyle birlikte yorumlanmalıdır. Dolayısıyla yüksek performans, bağımsız bir dış risk etiketinin tamamen keşfedildiğini değil, oluşturulan risk skorlama sisteminin öğrenilebilir ve içsel olarak tutarlı olduğunu göstermektedir.")
    add_heading(doc, "8.1 Öğrenilenler", 2)
    add_bullets(doc, [
        "EDA sonuçlarının hedef değişken seçimini doğrudan etkileyebileceği görüldü.",
        "Özellik mühendisliği ve eşik seçimi, model başarısı kadar metodolojik yorumlanabilirliği de belirler.",
        "Pipeline kullanımı ölçeklendirme işlemini eğitim akışının güvenli parçası haline getirir.",
        "Sınıf dengesizliği, accuracy tek başına yüksek olsa bile macro F1 yorumunu önemli kılar.",
        "VIF ve R² gibi kontroller, oluşturulan skor sisteminin istatistiksel tutarlılığını destekler.",
    ])
    add_heading(doc, "8.2 Gelecekte Yapılabilecek İyileştirmeler", 2)
    add_bullets(doc, [
        "GridSearchCV veya RandomizedSearchCV ile hiperparametre optimizasyonu yapılabilir.",
        "XGBoost ve LightGBM gibi gelişmiş topluluk yöntemleri denenebilir.",
        "Zaman boyutu güçlendirilmiş gerçek verilerle tedarikçi risk trendleri incelenebilir.",
        "Model sonuçları farklı sektörlerden dış veri setleriyle doğrulanabilir.",
        "Risk skor ağırlıkları uzman görüşü veya AHP/TOPSIS gibi MCDM yöntemleriyle kalibre edilebilir.",
    ])

    add_heading(doc, "9. Kaynakça", 1)
    refs = [
        "Baryannis, G., Dani, S., & Antoniou, G. (2019). Predicting supply chain risks using machine learning: The trade-off between performance and interpretability. Future Generation Computer Systems, 101, 993-1004. https://doi.org/10.1016/j.future.2019.07.059",
        "Ho, W., Xu, X., & Dey, P. K. (2010). Multi-criteria decision making approaches for supplier evaluation and selection: A literature review. European Journal of Operational Research, 202(1), 16-24. https://doi.org/10.1016/j.ejor.2009.05.009",
        "Tang, C. S. (2006). Perspectives in supply chain risk management. International Journal of Production Economics, 103(2), 451-488. https://doi.org/10.1016/j.ijpe.2005.12.006",
        "Kaggle. (2025). Supplier Risk Assessment Dataset. https://www.kaggle.com/datasets/programmer3/supplier-risk-assessment-dataset/data",
        "scikit-learn developers. (2026). Pipeline, StandardScaler and RandomForestClassifier documentation. https://scikit-learn.org/stable/",
        "pandas developers. (2026). pandas documentation. https://pandas.pydata.org/docs/",
        "NumPy developers. (2026). NumPy documentation. https://numpy.org/doc/stable/",
        "Matplotlib developers. (2026). Matplotlib documentation. https://matplotlib.org/stable/",
        "Waskom, M. (2026). seaborn documentation. https://seaborn.pydata.org/",
        "Wikimedia Commons. (2025). Yıldız Technical University Logo.png. https://commons.wikimedia.org/wiki/File:Yıldız_Technical_University_Logo.png",
    ]
    for ref in refs:
        add_para(doc, ref)

    add_heading(doc, "10. Ekler", 1)
    add_heading(doc, "10.1 Proje Dosyaları Envanteri", 2)
    add_table(doc, ["Dosya", "Açıklama"], appendix_inventory_rows(), widths=[3.8, 2.5], font_size=7.8)
    add_heading(doc, "10.2 Kod Ekleri", 2)
    for fn in ["eda_supplier_risk.py", "supplier_scoring.py", "supplier_scoring_quantile.py", "ml_modeling.py", "advanced_stats.py"]:
        add_code_appendix(doc, fn)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
